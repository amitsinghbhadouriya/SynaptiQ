"""Document ingestion: PDF (PyMuPDF), DOCX (python-docx) and TXT (standard library).

Each loader returns the cleaned full text plus section spans that remember where each
piece came from (page number for PDFs, heading for DOCX/TXT). Errors are raised as
DocumentLoadError with a stable code and a message that is safe to show users.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from pathlib import Path

from .schemas import LoadedDocument, SectionSpan
from .text_utils import normalize_whitespace

SUPPORTED_TYPES = {".pdf": "PDF", ".docx": "DOCX", ".txt": "TXT"}


class DocumentLoadError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def _safe_name(filename: str) -> str:
    name = Path(filename or "document").name
    name = re.sub(r"[^\w.\- ]+", "_", name).strip() or "document"
    return name[:120]


def _assemble(parts: list[tuple[str, int | None, str | None]]) -> tuple[str, list[SectionSpan]]:
    """Join (text, page, label) parts into one text, tracking each part's span."""
    pieces: list[str] = []
    sections: list[SectionSpan] = []
    cursor = 0
    for raw, page, label in parts:
        text = normalize_whitespace(raw)
        if not text:
            continue
        if pieces:
            pieces.append("\n\n")
            cursor += 2
        pieces.append(text)
        sections.append(SectionSpan(start=cursor, end=cursor + len(text), page_number=page, label=label))
        cursor += len(text)
    return "".join(pieces), sections


def _load_pdf(data: bytes) -> tuple[list[tuple[str, int | None, str | None]], dict]:
    if not data.startswith(b"%PDF"):
        raise DocumentLoadError("invalid_content", "The file does not look like a valid PDF.")
    try:
        import pymupdf as fitz  # PyMuPDF >= 1.24 module name
    except ImportError:  # pragma: no cover - older PyMuPDF releases
        try:
            import fitz
        except ImportError as exc:
            raise DocumentLoadError("missing_dependency", "PyMuPDF is required to read PDF files.") from exc
    try:
        with fitz.open(stream=data, filetype="pdf") as pdf:
            if pdf.needs_pass:
                raise DocumentLoadError("encrypted", "Password-protected PDFs are not supported.")
            parts = [(page.get_text("text"), index + 1, f"Page {index + 1}") for index, page in enumerate(pdf)]
            meta = {"pages": pdf.page_count, "title": (pdf.metadata or {}).get("title") or None}
    except DocumentLoadError:
        raise
    except Exception as exc:
        raise DocumentLoadError("parse_failed", "The PDF could not be parsed; it may be damaged.") from exc
    return parts, meta


def _load_docx(data: bytes) -> tuple[list[tuple[str, int | None, str | None]], dict]:
    if not data.startswith(b"PK") or not zipfile.is_zipfile(io.BytesIO(data)):
        raise DocumentLoadError("invalid_content", "The file does not look like a valid DOCX document.")
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentLoadError("missing_dependency", "python-docx is required to read DOCX files.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentLoadError("parse_failed", "The DOCX file could not be parsed; it may be damaged.") from exc

    parts: list[tuple[str, int | None, str | None]] = []
    heading = "Introduction"
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            parts.append(("\n".join(buffer), None, heading))
            buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.lower().startswith(("heading", "title")):
            flush()
            heading = text[:80]
            buffer.append(text)
        else:
            buffer.append(text)
    flush()
    for index, table in enumerate(document.tables):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        parts.append(("\n".join(r for r in rows if r.strip(" |")), None, f"Table {index + 1}"))
    return parts, {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}


def _load_txt(data: bytes) -> tuple[list[tuple[str, int | None, str | None]], dict]:
    if b"\x00" in data[:4096]:
        raise DocumentLoadError("invalid_content", "The file looks binary, not plain text.")
    for encoding in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 always decodes
        raise DocumentLoadError("decode_failed", "The text file encoding could not be detected.")

    # Treat markdown-style headings or ALL-CAPS lines as section boundaries.
    parts: list[tuple[str, int | None, str | None]] = []
    heading = "Document"
    buffer: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        is_heading = stripped.startswith("#") or (
            4 <= len(stripped) <= 80 and stripped.isupper() and any(c.isalpha() for c in stripped)
        )
        if is_heading:
            if any(b.strip() for b in buffer):
                parts.append(("\n".join(buffer), None, heading))
            buffer = []
            heading = stripped.lstrip("#").strip()[:80] or heading
        buffer.append(line.lstrip("#").strip() if is_heading else line)
    if any(b.strip() for b in buffer):
        parts.append(("\n".join(buffer), None, heading))
    return parts, {"encoding": encoding}


def load_document(filename: str, data: bytes, max_upload_mb: int = 20) -> LoadedDocument:
    """Validate and parse an uploaded file."""
    name = _safe_name(filename)
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_TYPES:
        raise DocumentLoadError(
            "unsupported_type", f"Unsupported file type '{suffix or 'none'}'. Upload a PDF, DOCX or TXT file."
        )
    if not data:
        raise DocumentLoadError("empty_file", "The file is empty.")
    if len(data) > max_upload_mb * 1024 * 1024:
        raise DocumentLoadError("too_large", f"The file is larger than the {max_upload_mb} MB limit.")

    file_type = SUPPORTED_TYPES[suffix]
    loader = {"PDF": _load_pdf, "DOCX": _load_docx, "TXT": _load_txt}[file_type]
    parts, meta = loader(data)
    text, sections = _assemble(parts)
    if not text.strip():
        hint = " Scanned PDFs need OCR, which is outside the current scope." if file_type == "PDF" else ""
        raise DocumentLoadError("no_text", f"No extractable text was found in the file.{hint}")

    meta.update({"sections": [s.__dict__ for s in sections], "section_count": len(sections), "bytes": len(data)})
    if file_type == "PDF":
        meta["pages_with_text"] = len({s.page_number for s in sections})
    return LoadedDocument(
        filename=name,
        file_type=file_type,
        content_hash=hashlib.sha256(data).hexdigest(),
        text=text,
        sections=sections,
        metadata=meta,
    )
